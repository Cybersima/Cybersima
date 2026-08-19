from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO, StringIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from securetrade.engine.market_colors import fill_ratio_for, outcome_tone
from securetrade.models import Opportunity, PaperPosition

HEADERS = [
    "ID",
    "Detected Time",
    "Strategy",
    "Market",
    "Route",
    "Financial",
    "Execution",
    "Guardian",
    "Guardian Score",
    "Expected Profit",
    "Realized Profit",
    "Edge Lifetime",
    "Close Reason",
    "Buy Venue",
    "Sell Venue",
    "Raw Edge",
    "Net Edge (bps)",
    "Fee (bps)",
    "Slippage (bps)",
    "Fill Ratio",
    "Paper Notional",
]

MARKET_COL = HEADERS.index("Market") + 1

# Same tones as the command-center market tiles.
TONE_FILLS = {
    "profit": PatternFill("solid", fgColor="3DD68C"),
    "loss": PatternFill("solid", fgColor="FF6B6B"),
    "missed": PatternFill("solid", fgColor="6CB6FF"),
    "reversal": PatternFill("solid", fgColor="F5A14A"),
}

STRATEGY_LABELS = {
    "cross_venue": "Cross-venue",
    "triangular": "Triangular",
    "alert": "Alert",
    "cex_dex": "CEX–DEX",
}


@dataclass
class ProfitReportRow:
    id: str
    detected_time: str
    strategy: str
    market: str
    route: str
    financial: str
    execution: str
    guardian: str
    guardian_score: int
    expected_profit: float
    realized_profit: float
    edge_lifetime: float
    close_reason: str
    buy_venue: str
    sell_venue: str
    raw_edge: float
    net_edge_bps: float
    fee_bps: float
    slippage_bps: float
    fill_ratio: float
    paper_notional: float
    tone: str = ""
    extra: dict[str, Any] = field(default_factory=dict)

    def values(self) -> list[Any]:
        return [
            self.id,
            self.detected_time,
            self.strategy,
            self.market,
            self.route,
            self.financial,
            self.execution,
            self.guardian,
            self.guardian_score,
            self.expected_profit,
            self.realized_profit,
            self.edge_lifetime,
            self.close_reason,
            self.buy_venue,
            self.sell_venue,
            self.raw_edge,
            self.net_edge_bps,
            self.fee_bps,
            self.slippage_bps,
            self.fill_ratio,
            self.paper_notional,
        ]

    def to_dict(self) -> dict[str, Any]:
        data = dict(zip(HEADERS, self.values()))
        data["tone"] = self.tone
        return data


def fill_ratio_for(outcome: str) -> float:
    flag = (outcome or "").upper()
    if flag == "CAPTURED":
        return 1.0
    if flag == "REVERSED":
        return 0.5
    return 0.0


def format_detected(ts: float) -> str:
    if not ts:
        return ""
    return datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def strategy_label(kind: Any) -> str:
    key = kind.value if hasattr(kind, "value") else str(kind or "")
    return STRATEGY_LABELS.get(key, key.replace("_", " ").title() or "—")


def legs_venues(opportunity: Opportunity) -> tuple[str, str]:
    buys = [leg.venue for leg in opportunity.legs if (leg.action or "").lower() == "buy"]
    sells = [leg.venue for leg in opportunity.legs if (leg.action or "").lower() == "sell"]
    buy = buys[0] if buys else (opportunity.legs[0].venue if opportunity.legs else "")
    sell = sells[0] if sells else (opportunity.legs[1].venue if len(opportunity.legs) > 1 else "")
    return buy, sell


def position_extras(opportunity: Opportunity, financial: str = "Paper") -> dict[str, Any]:
    buy, sell = legs_venues(opportunity)
    return {
        "detected_at": opportunity.ts,
        "strategy": strategy_label(opportunity.kind),
        "route": f"{buy} → {sell}" if sell else buy,
        "financial": financial,
        "guardian": "ALLOW" if opportunity.guardian_allowed else "BLOCK",
        "buy_venue": buy,
        "sell_venue": sell,
        "raw_edge_bps": round(opportunity.edge_bps, 4),
        "fee_bps": round(opportunity.estimated_fees_bps, 4),
        "slippage_bps": round(opportunity.estimated_slippage_bps, 4),
    }


def _round(value: float, digits: int = 4) -> float:
    return round(float(value or 0.0), digits)


def row_from_position(position: PaperPosition, now: float | None = None) -> ProfitReportRow:
    start = position.detected_at or position.opened_at
    end = position.closed_at or now or start
    lifetime = max(0.0, float(end) - float(start or 0.0))
    reason = position.outcome or ""
    pnl = float(position.actual_pnl or 0.0)
    ratio = position.fill_ratio or fill_ratio_for(reason)
    return ProfitReportRow(
        id=position.opportunity_id,
        detected_time=format_detected(start),
        strategy=position.strategy or "—",
        market=position.pair,
        route=position.route or _route(position.buy_venue, position.sell_venue),
        financial=position.financial or "Paper",
        execution=position.commit_kind or "—",
        guardian=position.guardian or "ALLOW",
        guardian_score=int(position.trust_score or 0),
        expected_profit=_round(position.expected_pnl),
        realized_profit=_round(pnl),
        edge_lifetime=_round(lifetime, 3),
        close_reason=reason,
        buy_venue=position.buy_venue,
        sell_venue=position.sell_venue,
        raw_edge=_round(position.raw_edge_bps),
        net_edge_bps=_round(position.expected_net_edge_bps or position.actual_net_edge_bps),
        fee_bps=_round(position.fee_bps),
        slippage_bps=_round(position.slippage_bps),
        fill_ratio=_round(ratio, 2),
        paper_notional=_round(position.notional, 2),
        tone=outcome_tone(reason, pnl),
    )


def row_from_opportunity(
    opportunity: Opportunity,
    *,
    financial: str,
    execution: str,
    close_reason: str,
    realized_profit: float = 0.0,
    expected_profit: float | None = None,
    edge_lifetime_s: float = 0.0,
    net_edge_bps: float | None = None,
) -> ProfitReportRow:
    extras = position_extras(opportunity, financial)
    buy, sell = extras["buy_venue"], extras["sell_venue"]
    net = net_edge_bps if net_edge_bps is not None else (opportunity.expected_net_edge_bps or opportunity.net_edge_bps)
    expected = expected_profit if expected_profit is not None else opportunity.notional * (net / 10_000)
    pnl = float(realized_profit or 0.0)
    return ProfitReportRow(
        id=opportunity.id,
        detected_time=format_detected(opportunity.ts),
        strategy=extras["strategy"],
        market=opportunity.pair or extras["route"],
        route=extras["route"],
        financial=financial,
        execution=execution,
        guardian="ALLOW" if opportunity.guardian_allowed else "BLOCK",
        guardian_score=int(opportunity.trust_score or 0),
        expected_profit=_round(expected),
        realized_profit=_round(pnl),
        edge_lifetime=_round(edge_lifetime_s, 3),
        close_reason=close_reason,
        buy_venue=buy,
        sell_venue=sell,
        raw_edge=_round(opportunity.edge_bps),
        net_edge_bps=_round(net),
        fee_bps=_round(opportunity.estimated_fees_bps),
        slippage_bps=_round(opportunity.estimated_slippage_bps),
        fill_ratio=fill_ratio_for(close_reason),
        paper_notional=_round(opportunity.notional, 2),
        tone=outcome_tone(close_reason, pnl),
    )


def _route(buy: str, sell: str) -> str:
    if buy and sell:
        return f"{buy} → {sell}"
    return buy or sell or ""


def rows_to_csv(rows: Iterable[ProfitReportRow]) -> str:
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    for row in rows:
        writer.writerow(row.values())
    return buf.getvalue()


def rows_to_xlsx(rows: Iterable[ProfitReportRow], title: str = "Profit Report") -> bytes:
    workbook = Workbook()
    sheet: Worksheet = workbook.active
    sheet.title = title[:31]
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF6")
    thin = Border(
        left=Side(style="thin", color="C0C0C0"),
        right=Side(style="thin", color="C0C0C0"),
        top=Side(style="thin", color="C0C0C0"),
        bottom=Side(style="thin", color="C0C0C0"),
    )
    for col, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin
    money_cols = {
        HEADERS.index("Expected Profit") + 1,
        HEADERS.index("Realized Profit") + 1,
        HEADERS.index("Paper Notional") + 1,
    }
    number_cols = {
        HEADERS.index("Guardian Score") + 1,
        HEADERS.index("Edge Lifetime") + 1,
        HEADERS.index("Raw Edge") + 1,
        HEADERS.index("Net Edge (bps)") + 1,
        HEADERS.index("Fee (bps)") + 1,
        HEADERS.index("Slippage (bps)") + 1,
        HEADERS.index("Fill Ratio") + 1,
    }
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row.values(), start=1):
            cell = sheet.cell(r_idx, c_idx, value)
            cell.border = thin
            if c_idx in money_cols:
                cell.number_format = "0.00"
            elif c_idx in number_cols:
                cell.number_format = "0.00"
        market = sheet.cell(r_idx, MARKET_COL)
        fill = TONE_FILLS.get(row.tone)
        if fill:
            market.fill = fill
            market.font = Font(bold=True)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, sheet.max_row)}"
    sheet.freeze_panes = "A2"
    widths = {
        "ID": 28,
        "Detected Time": 24,
        "Strategy": 14,
        "Market": 14,
        "Route": 28,
        "Financial": 12,
        "Execution": 18,
        "Guardian": 12,
        "Guardian Score": 16,
        "Expected Profit": 16,
        "Realized Profit": 16,
        "Edge Lifetime": 14,
        "Close Reason": 14,
        "Buy Venue": 14,
        "Sell Venue": 14,
        "Raw Edge": 12,
        "Net Edge (bps)": 16,
        "Fee (bps)": 12,
        "Slippage (bps)": 16,
        "Fill Ratio": 12,
        "Paper Notional": 16,
    }
    for col, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = widths.get(header, 14)
    buf = BytesIO()
    workbook.save(buf)
    return buf.getvalue()
