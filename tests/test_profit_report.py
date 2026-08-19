from io import BytesIO

import pytest
from openpyxl import load_workbook

from securetrade.config import AppConfig
from securetrade.engine.market_colors import fill_ratio_for, outcome_tone
from securetrade.engine.profit_report import HEADERS, rows_to_xlsx
from securetrade.engine.runner import Engine
from securetrade.models import CommitDecision, OperatingMode, PaperOutcome


@pytest.fixture
def engine() -> Engine:
    config = AppConfig()
    config.env.demo_only = True
    config._apply_env_overrides()
    eng = Engine(config)
    eng.set_mode(OperatingMode.LEARN)
    return eng


@pytest.mark.asyncio
async def test_profit_report_extracts_headers_and_trade_fields(engine: Engine) -> None:
    captured = await engine.force_outcome("captured")
    reversed_pos = await engine.force_outcome("reversed")
    cancelled = await engine.force_outcome("cancel")
    missed = await engine.force_outcome("missed")

    rows = {row.id: row for row in engine.profit_rows()}
    cap = rows[captured.opportunity_id]
    rev = rows[reversed_pos.opportunity_id]
    can = rows[cancelled.opportunity_id]
    miss = rows[missed.opportunity_id]

    assert cap.strategy == "Cross-venue"
    assert cap.market == "BTC/USDC"
    assert cap.buy_venue == "coinbase"
    assert cap.sell_venue == "kraken"
    assert "coinbase" in cap.route and "kraken" in cap.route
    assert cap.financial == "Paper"
    assert cap.execution in {CommitDecision.COMMIT.value, CommitDecision.RESEARCH_COMMIT.value}
    assert cap.guardian == "ALLOW"
    assert cap.guardian_score >= 0
    assert cap.expected_profit > 0
    assert cap.realized_profit == pytest.approx(0.31)
    assert cap.close_reason == PaperOutcome.CAPTURED.value
    assert cap.raw_edge > 0
    assert cap.net_edge_bps > 0
    assert cap.fill_ratio == 1.0
    assert cap.paper_notional == 100
    assert cap.tone == "profit"
    assert cap.detected_time.endswith("UTC")

    assert rev.close_reason == PaperOutcome.REVERSED.value
    assert rev.realized_profit < 0
    assert rev.fill_ratio == 0.5
    assert rev.tone == "reversal"

    assert can.execution == CommitDecision.CANCEL.value
    assert can.close_reason == "CANCEL"
    assert can.fill_ratio == 0.0
    assert can.tone == "missed"

    assert miss.close_reason == PaperOutcome.MISSED.value
    assert miss.fill_ratio == 0.0
    assert miss.tone == "missed"

    payload = engine.profit_report_csv()
    first = payload.splitlines()[0]
    assert first.split(",") == HEADERS
    assert captured.opportunity_id in payload
    assert "BTC/USDC" in payload
    assert "coinbase" in payload


@pytest.mark.asyncio
async def test_excel_colors_market_column(engine: Engine) -> None:
    await engine.force_outcome("captured")
    await engine.force_outcome("loss")
    await engine.force_outcome("missed")
    await engine.force_outcome("reversed")

    data = engine.profit_report_xlsx()
    workbook = load_workbook(BytesIO(data))
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == HEADERS

    fills: list[tuple[str, float, str]] = []
    for row in sheet.iter_rows(min_row=2, max_col=len(HEADERS)):
        reason = str(row[HEADERS.index("Close Reason")].value)
        realized = float(row[HEADERS.index("Realized Profit")].value or 0)
        market = row[HEADERS.index("Market")]
        rgb = str(market.fill.fgColor.rgb).upper()
        fills.append((reason, realized, rgb))
        assert market.value == "BTC/USDC"

    assert any(reason == "CAPTURED" and pnl > 0 and rgb.endswith("3DD68C") for reason, pnl, rgb in fills)
    assert any(reason == "CAPTURED" and pnl < 0 and rgb.endswith("FF6B6B") for reason, pnl, rgb in fills)
    assert any(reason == "REVERSED" and rgb.endswith("F5A14A") for reason, pnl, rgb in fills)
    assert any(reason == "MISSED" and rgb.endswith("6CB6FF") for reason, pnl, rgb in fills)


def test_empty_workbook_still_has_headers() -> None:
    workbook = load_workbook(BytesIO(rows_to_xlsx([])))
    assert [cell.value for cell in workbook.active[1]] == HEADERS


def test_fill_ratio_and_cancelled_tone() -> None:
    assert fill_ratio_for("CAPTURED") == 1.0
    assert fill_ratio_for("REVERSED") == 0.5
    assert fill_ratio_for("CANCEL") == 0.0
    assert outcome_tone("CANCELLED") == "missed"
    assert outcome_tone("CANCEL") == "missed"
    assert outcome_tone("BLOCKED") == "missed"
